#!/usr/bin/env python3
"""Build a single non-technical stakeholder briefing for gitlab-compliance."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = RGBColor(0x10, 0x22, 0x30)
SEA = RGBColor(0x1A, 0x6B, 0x63)
SEA_SOFT = RGBColor(0x3C, 0xA0, 0x94)
AMBER = RGBColor(0xC4, 0x7A, 0x28)
PAPER = RGBColor(0xF5, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5E, 0x70, 0x7A)
BODY = RGBColor(0x2A, 0x3B, 0x46)

TOTAL = 12


def _font(run, *, size: Pt, bold: bool = False, color: RGBColor = BODY, name: str = "Calibri") -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.insert(0, latin)
    latin.set("typeface", name)


def _rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _box(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _text(shape, text: str, *, size: Pt, bold: bool = False, color: RGBColor = BODY, align=PP_ALIGN.LEFT) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        _font(run, size=size, bold=bold, color=color)


def _bullets(shape, items: list[str], *, size: Pt = Pt(16), color: RGBColor = BODY) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.space_after = Pt(9)
        for run in p.runs:
            _font(run, size=size, color=color)


def _lines(shape, lines: list[tuple[str, RGBColor, bool, Pt]]) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (text, color, bold, size) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(5)
        for run in p.runs:
            _font(run, size=size, bold=bold, color=color)


def _header(slide, title: str, subtitle: str | None = None) -> None:
    _rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.25), INK)
    _rect(slide, Inches(0), Inches(1.25), SLIDE_W, Inches(0.07), SEA)
    _text(_box(slide, Inches(0.65), Inches(0.25), Inches(12), Inches(0.5)), title, size=Pt(24), bold=True, color=WHITE)
    if subtitle:
        _text(_box(slide, Inches(0.65), Inches(0.78), Inches(12), Inches(0.35)), subtitle, size=Pt(13), color=SEA_SOFT)


def _footer(slide, page: int) -> None:
    _text(
        _box(slide, Inches(0.65), Inches(7.05), Inches(9), Inches(0.28)),
        "Gitlab Compliance  ·  Stakeholder briefing",
        size=Pt(11),
        color=MUTED,
    )
    _text(
        _box(slide, Inches(11.2), Inches(7.05), Inches(1.5), Inches(0.28)),
        f"{page} / {TOTAL}",
        size=Pt(11),
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _bg(slide) -> None:
    _rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, PAPER)


def _card(slide, left, top, width, height):
    return _rect(slide, left, top, width, height, WHITE)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    _rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, INK)
    _rect(s, Inches(0), Inches(0), Inches(0.2), SLIDE_H, SEA)
    _rect(s, Inches(0), Inches(5.85), SLIDE_W, Inches(1.65), RGBColor(0x17, 0x30, 0x3F))
    _text(
        _box(s, Inches(0.85), Inches(1.5), Inches(11.5), Inches(0.35)),
        "STAKEHOLDER BRIEFING",
        size=Pt(13),
        bold=True,
        color=SEA_SOFT,
    )
    _text(
        _box(s, Inches(0.85), Inches(2.0), Inches(11.5), Inches(1.4)),
        "Gitlab Compliance",
        size=Pt(42),
        bold=True,
        color=WHITE,
    )
    _text(
        _box(s, Inches(0.85), Inches(3.4), Inches(11), Inches(1.2)),
        "Find delivery risk early, meet Nexus expectations, support the GitLab Dedicated move — "
        "and keep standards clear for everyone afterwards.",
        size=Pt(18),
        color=RGBColor(0xC0, 0xD0, 0xD8),
    )
    _lines(
        _box(s, Inches(0.85), Inches(6.2), Inches(11.5), Inches(0.9)),
        [
            ("For delivery, security, and programme leaders", WHITE, True, Pt(15)),
            ("Plain-language rules · clearer readiness · less last-minute surprise", SEA_SOFT, False, Pt(13)),
        ],
    )

    # 2 Agenda
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "What we will cover", "Outcomes first — not tooling detail")
    _footer(s, 2)
    _bullets(
        _box(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(4.8)),
        [
            "1.  The problem we are solving",
            "2.  Why it matters for Dedicated and Nexus",
            "3.  What Gitlab Compliance does in plain terms",
            "4.  How we find and fix delivery risk",
            "5.  How easy-to-read policies keep standards ongoing",
            "6.  What good looks like — and the ask",
        ],
        size=Pt(20),
    )

    # 3 The problem
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "The problem", "Important delivery rules are easy to miss until something breaks")
    _footer(s, 3)
    cards = [
        ("Hidden risk", "Build and release setups can quietly drift from agreed standards."),
        ("Late discovery", "Issues often show up close to go-live — when change is costly."),
        ("Hard to explain", "Rules live in long technical files that few people can review with confidence."),
        ("Uneven practice", "Some teams follow the rules; others do not — and leaders cannot see the gap."),
    ]
    for i, (h, body) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.65 + col * 6.3)
        top = Inches(1.7 + row * 2.4)
        _card(s, left, top, Inches(6.0), Inches(2.15))
        _rect(s, left, top, Inches(0.12), Inches(2.15), AMBER if i % 2 else SEA)
        _text(_box(s, left + Inches(0.4), top + Inches(0.35), Inches(5.3), Inches(0.45)), h, size=Pt(18), bold=True, color=INK)
        _text(_box(s, left + Inches(0.4), top + Inches(0.95), Inches(5.3), Inches(0.9)), body, size=Pt(15), color=BODY)

    # 4 Why now
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Why this matters now", "Two programmes that make invisible risk visible")
    _footer(s, 4)
    _card(s, Inches(0.65), Inches(1.65), Inches(5.9), Inches(4.7))
    _rect(s, Inches(0.65), Inches(1.65), Inches(5.9), Inches(0.7), SEA)
    _text(_box(s, Inches(0.9), Inches(1.8), Inches(5.4), Inches(0.45)), "GitLab Dedicated migration", size=Pt(17), bold=True, color=WHITE)
    _bullets(
        _box(s, Inches(0.95), Inches(2.6), Inches(5.3), Inches(3.4)),
        [
            "We need a clear picture of delivery readiness before cutover",
            "Unknown debt turns into schedule risk",
            "Leaders need simple evidence — not technical archaeology",
            "Better to fix early than discover during migration",
        ],
        size=Pt(15),
    )
    _card(s, Inches(6.8), Inches(1.65), Inches(5.9), Inches(4.7))
    _rect(s, Inches(6.8), Inches(1.65), Inches(5.9), Inches(0.7), AMBER)
    _text(_box(s, Inches(7.05), Inches(1.8), Inches(5.4), Inches(0.45)), "Nexus compliance demand", size=Pt(17), bold=True, color=WHITE)
    _bullets(
        _box(s, Inches(7.1), Inches(2.6), Inches(5.3), Inches(3.4)),
        [
            "Approved, trusted components become mandatory",
            "Loose or outdated dependencies create control gaps",
            "We need proof that standards are followed",
            "Fixes must be reviewable and sustainable",
        ],
        size=Pt(15),
    )

    # 5 What it is
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "What Gitlab Compliance does", "A practical way to check delivery setups against agreed standards")
    _footer(s, 5)
    _text(
        _box(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.55)),
        "It checks how teams build and release software — before changes are approved — and highlights where practice does not match the standard.",
        size=Pt(16),
        color=BODY,
    )
    points = [
        ("Find issues early", "Spot delivery risk while there is still time to fix it."),
        ("Speak plainly", "Standards are written in everyday language teams can share."),
        ("Show the evidence", "Clear reports for programme, security, and delivery forums."),
        ("Help fix what matters", "Guided remediations reduce the backlog of known debt."),
        ("Keep knowledge current", "Living summaries of how pipelines actually work."),
        ("Stay consistent", "The same standard applies across teams and projects."),
    ]
    for i, (h, body) in enumerate(points):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(2.3 + row * 2.2)
        _card(s, left, top, Inches(4.0), Inches(2.0))
        _rect(s, left, top, Inches(4.0), Inches(0.1), SEA)
        _text(_box(s, left + Inches(0.25), top + Inches(0.3), Inches(3.5), Inches(0.45)), h, size=Pt(15), bold=True, color=INK)
        _text(_box(s, left + Inches(0.25), top + Inches(0.9), Inches(3.5), Inches(0.85)), body, size=Pt(13), color=BODY)

    # 6 How we work
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "How we use it", "Four simple activities — one outcome: fewer surprises")
    _footer(s, 6)
    caps = [
        ("Check", "Compare each project against the agreed standard and list what needs attention."),
        ("Scan widely", "Review many projects quickly so the migration cohort has a ranked risk view."),
        ("Strengthen supply chain", "Tighten use of trusted components and close known gaps against Nexus expectations."),
        ("Document clearly", "Produce readable summaries so non-specialists can join the readiness conversation."),
    ]
    for i, (h, body) in enumerate(caps):
        top = Inches(1.55 + i * 1.2)
        _card(s, Inches(0.65), top, Inches(12), Inches(1.08))
        _rect(s, Inches(0.65), top, Inches(2.6), Inches(1.08), SEA)
        _text(
            _box(s, Inches(0.75), top + Inches(0.32), Inches(2.4), Inches(0.45)),
            h,
            size=Pt(15),
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        _text(_box(s, Inches(3.5), top + Inches(0.3), Inches(8.8), Inches(0.55)), body, size=Pt(15), color=BODY)

    # 7 Tech debt in business terms
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "The risk we need to surface", "Common delivery debt — explained without jargon")
    _footer(s, 7)
    debts = [
        ("Untrusted building blocks", "Projects using components that can change without notice."),
        ("Out-of-date dependencies", "Critical parts lagging behind agreed freshness windows."),
        ("Weak release guards", "Steps that can run when they should not."),
        ("Bypassed standards", "Teams skipping the shared delivery patterns."),
        ("Poor visibility", "No clear picture of how a pipeline works today."),
        ("Policy only on paper", "Rules exist in guidance, but are not checked automatically."),
    ]
    for i, (h, body) in enumerate(debts):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.65 + row * 2.4)
        _card(s, left, top, Inches(4.0), Inches(2.15))
        _text(_box(s, left + Inches(0.25), top + Inches(0.35), Inches(3.5), Inches(0.55)), h, size=Pt(15), bold=True, color=SEA)
        _text(_box(s, left + Inches(0.25), top + Inches(1.0), Inches(3.5), Inches(0.85)), body, size=Pt(13), color=BODY)

    # 8 Journey
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "The journey", "From unknown risk to confident, ongoing control")
    _footer(s, 8)
    steps = [
        ("1", "Discover", "Scan the migration cohort and produce a clear risk list."),
        ("2", "Prioritise", "Focus first on the highest-impact projects and gaps."),
        ("3", "Remediate", "Fix supply-chain and standards issues with reviewable changes."),
        ("4", "Explain", "Share readable reports and pipeline summaries with stakeholders."),
        ("5", "Maintain", "Keep the same plain-language standards in force after go-live."),
    ]
    for i, (num, title, body) in enumerate(steps):
        top = Inches(1.55 + i * 0.95)
        _card(s, Inches(0.65), top, Inches(12), Inches(0.85))
        _rect(s, Inches(0.8), top + Inches(0.15), Inches(0.7), Inches(0.55), SEA)
        _text(
            _box(s, Inches(0.8), top + Inches(0.25), Inches(0.7), Inches(0.4)),
            num,
            size=Pt(16),
            bold=True,
            color=WHITE,
            align=PP_ALIGN.CENTER,
        )
        _text(_box(s, Inches(1.8), top + Inches(0.12), Inches(3.2), Inches(0.35)), title, size=Pt(16), bold=True, color=INK)
        _text(_box(s, Inches(5.1), top + Inches(0.2), Inches(7.2), Inches(0.5)), body, size=Pt(14), color=BODY)

    # 9 Readable policies
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Standards people can understand", "Easy-to-read policies keep the rules alive after migration")
    _footer(s, 9)
    _text(
        _box(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.5)),
        "Migration finds the debt once. Clear policies make sure the same expectations stay true every day after.",
        size=Pt(15),
        color=BODY,
    )
    _card(s, Inches(0.65), Inches(2.15), Inches(5.9), Inches(4.2))
    _text(_box(s, Inches(0.95), Inches(2.4), Inches(5.3), Inches(0.4)), "Business benefits", size=Pt(16), bold=True, color=INK)
    _bullets(
        _box(s, Inches(0.95), Inches(3.0), Inches(5.3), Inches(3.0)),
        [
            "Security and delivery share one understandable contract",
            "New Nexus or Dedicated rules can be written plainly",
            "Ownership is clear — standards are not buried in tribal knowledge",
            "Findings are explainable in programme forums",
            "Standards survive team and supplier change",
        ],
        size=Pt(14),
    )
    _card(s, Inches(6.8), Inches(2.15), Inches(5.9), Inches(4.2))
    _text(_box(s, Inches(7.1), Inches(2.4), Inches(5.3), Inches(0.4)), "Example of a readable rule", size=Pt(16), bold=True, color=INK)
    _text(
        _box(s, Inches(7.1), Inches(3.1), Inches(5.3), Inches(2.8)),
        "“If a job uses a software image, it must not use an open-ended ‘latest’ version.”\n\n"
        "That idea becomes a short scenario both technical and non-technical colleagues can follow — "
        "and the check confirms it is true before approval.",
        size=Pt(14),
        color=BODY,
    )

    # 10 Value for Dedicated + Nexus
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "How this supports our programmes", "One approach serving migration and compliance together")
    _footer(s, 10)
    rows = [
        ("Need", "How Gitlab Compliance helps"),
        ("Know readiness before Dedicated cutover", "Ranked risk view across the migration cohort"),
        ("Meet Nexus expectations", "Checks and fixes for trusted, up-to-date components"),
        ("Reduce last-minute surprise", "Issues found and discussed before go-live pressure"),
        ("Give leaders usable evidence", "Clear reports and summaries for decision forums"),
        ("Keep standards after the move", "Readable policies remain the ongoing contract"),
    ]
    for i, (left, right) in enumerate(rows):
        top = Inches(1.5 + i * 0.8)
        bg = SEA if i == 0 else WHITE
        lc = WHITE if i == 0 else INK
        rc = WHITE if i == 0 else BODY
        _rect(s, Inches(0.65), top, Inches(5.7), Inches(0.7), bg)
        _rect(s, Inches(6.45), top, Inches(6.2), Inches(0.7), bg)
        _text(_box(s, Inches(0.85), top + Inches(0.18), Inches(5.3), Inches(0.4)), left, size=Pt(13), bold=(i == 0), color=lc)
        _text(_box(s, Inches(6.65), top + Inches(0.18), Inches(5.8), Inches(0.4)), right, size=Pt(13), bold=(i == 0), color=rc)

    # 11 What good looks like
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "What good looks like", "Simple measures for programme confidence")
    _footer(s, 11)
    measures = [
        ("Coverage", "Most migration projects have been reviewed against the standard."),
        ("Open issues falling", "Critical gaps are tracked and trending down."),
        ("Faster fixes", "Time from finding to approved fix is shortening."),
        ("Clear explanations", "Teams can explain a finding without specialist decoding."),
        ("Ready evidence", "Readiness packs include current summaries and reports."),
        ("Standards stick", "The same rules still protect delivery after cutover."),
    ]
    for i, (h, body) in enumerate(measures):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.65 + row * 2.4)
        _card(s, left, top, Inches(4.0), Inches(2.15))
        _text(_box(s, left + Inches(0.25), top + Inches(0.35), Inches(3.5), Inches(0.45)), h, size=Pt(16), bold=True, color=SEA)
        _text(_box(s, left + Inches(0.25), top + Inches(0.95), Inches(3.5), Inches(0.9)), body, size=Pt(13), color=BODY)

    # 12 Ask / next steps
    s = prs.slides.add_slide(blank)
    _rect(s, Inches(0), Inches(0), SLIDE_W, SLIDE_H, INK)
    _rect(s, Inches(0), Inches(0), Inches(0.2), SLIDE_H, SEA)
    _text(_box(s, Inches(0.85), Inches(1.2), Inches(11.5), Inches(0.7)), "The ask", size=Pt(34), bold=True, color=WHITE)
    actions = [
        "Endorse a readiness scan of the Dedicated migration cohort",
        "Agree that Nexus-related standards will be written in plain language and checked before approval",
        "Support a warn-first approach, then strengthen once the baseline is clean",
        "Use the reports in programme forums as the shared view of risk and progress",
        "Confirm that the same readable standards remain in force after cutover",
    ]
    box = _box(s, Inches(0.85), Inches(2.15), Inches(11.5), Inches(3.6))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(actions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{i + 1}.  {item}"
        p.space_after = Pt(12)
        for run in p.runs:
            _font(run, size=Pt(16), color=WHITE)
    _text(
        _box(s, Inches(0.85), Inches(6.2), Inches(11.5), Inches(0.5)),
        "Outcome: fewer surprises, clearer ownership, and standards that stay understandable.",
        size=Pt(14),
        color=SEA_SOFT,
    )

    out = Path("/workspace/docs/presentations/gitlab-compliance-stakeholder-briefing.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    print(f"Wrote {build()}")

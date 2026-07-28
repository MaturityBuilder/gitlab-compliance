#!/usr/bin/env python3
"""Build a product overview PowerPoint from gitlab-compliance documentation."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Calm slate + sea-green direction (distinct from purple/cream clichés)
INK = RGBColor(0x12, 0x24, 0x32)
SEA = RGBColor(0x1A, 0x6B, 0x63)
SEA_SOFT = RGBColor(0x3C, 0xA0, 0x94)
AMBER = RGBColor(0xC9, 0x7B, 0x2A)
PAPER = RGBColor(0xF5, 0xF7, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5E, 0x70, 0x7A)
BODY = RGBColor(0x2A, 0x3B, 0x46)
SOFT_LINE = RGBColor(0xD7, 0xE0, 0xE4)

TOTAL = 14


def _font(run, *, size: Pt, bold: bool = False, color: RGBColor = BODY, name: str = "Calibri") -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.insert(0, latin)
    latin.set("typeface", name)


def _rect(slide, left, top, width, height, fill: RGBColor):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _box(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _text(shape, text: str, *, size: Pt, bold: bool = False, color: RGBColor = BODY, align=PP_ALIGN.LEFT) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        _font(run, size=size, bold=bold, color=color)


def _bullets(shape, items: list[str], *, size: Pt = Pt(17), color: RGBColor = BODY) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.space_after = Pt(9)
        for run in p.runs:
            _font(run, size=size, color=color)


def _lines(shape, lines: list[tuple[str, RGBColor, bool, Pt]], *, mono: bool = False) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    name = "Consolas" if mono else "Calibri"
    for i, (text, color, bold, size) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.space_after = Pt(4)
        for run in p.runs:
            _font(run, size=size, bold=bold, color=color, name=name)


def _header(slide, title: str, subtitle: str | None = None) -> None:
    _rect(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.28), INK)
    _rect(slide, Inches(0), Inches(1.28), SLIDE_WIDTH, Inches(0.07), SEA)
    box = _box(slide, Inches(0.65), Inches(0.28), Inches(12), Inches(0.55))
    _text(box, title, size=Pt(26), bold=True, color=WHITE)
    if subtitle:
        sub = _box(slide, Inches(0.65), Inches(0.8), Inches(12), Inches(0.35))
        _text(sub, subtitle, size=Pt(13), color=SEA_SOFT)


def _footer(slide, page: int) -> None:
    brand = _box(slide, Inches(0.65), Inches(7.05), Inches(9), Inches(0.28))
    _text(brand, "Based on product documentation  ·  gitlab-compliance", size=Pt(11), color=MUTED)
    num = _box(slide, Inches(11.2), Inches(7.05), Inches(1.5), Inches(0.28))
    _text(num, f"{page} / {TOTAL}", size=Pt(11), color=MUTED, align=PP_ALIGN.RIGHT)


def _bg(slide) -> None:
    _rect(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, PAPER)


def _card(slide, left, top, width, height):
    return _rect(slide, left, top, width, height, WHITE)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank = prs.slide_layouts[6]

    # 1. Title
    s = prs.slides.add_slide(blank)
    _rect(s, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, INK)
    _rect(s, Inches(0), Inches(0), Inches(0.2), SLIDE_HEIGHT, SEA)
    _rect(s, Inches(0), Inches(5.9), SLIDE_WIDTH, Inches(1.6), RGBColor(0x18, 0x32, 0x42))
    eye = _box(s, Inches(0.85), Inches(1.55), Inches(11), Inches(0.35))
    _text(eye, "PRODUCT OVERVIEW  ·  FROM THE DOCUMENTATION", size=Pt(13), bold=True, color=SEA_SOFT)
    title = _box(s, Inches(0.85), Inches(2.05), Inches(11.5), Inches(1.1))
    _text(title, "Gitlab Compliance", size=Pt(44), bold=True, color=WHITE)
    sub = _box(s, Inches(0.85), Inches(3.25), Inches(11), Inches(1.2))
    _text(
        sub,
        "BDD compliance testing and pipeline documentation for GitLab CI/CD — "
        "catch misconfigurations before merge, in language teams can share.",
        size=Pt(18),
        color=RGBColor(0xC2, 0xD2, 0xDA),
    )
    meta = _box(s, Inches(0.85), Inches(6.2), Inches(11.5), Inches(0.9))
    _lines(
        meta,
        [
            ("MaturityBuilder", WHITE, True, Pt(16)),
            (
                "Docs: maturitybuilder.github.io/gitlab-compliance  ·  PyPI: gitlab-compliance",
                SEA_SOFT,
                False,
                Pt(13),
            ),
        ],
    )

    # 2. Agenda
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Agenda", "What the documentation covers")
    _footer(s, 2)
    agenda = [
        "1.  Problem and product idea",
        "2.  Two workflows: check and generate",
        "3.  How BDD policies work",
        "4.  Starter security policy pack",
        "5.  Install, reports, and documentation output",
        "6.  CI/CD, OCI packs, and org-wide rollout",
        "7.  Where to go next",
    ]
    box = _box(s, Inches(0.9), Inches(1.8), Inches(11), Inches(4.8))
    _bullets(box, agenda, size=Pt(20))

    # 3. Problem / idea
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "The idea", "Negative testing for GitLab CI configuration")
    _footer(s, 3)
    lead = _box(s, Inches(0.7), Inches(1.65), Inches(12), Inches(0.7))
    _text(
        lead,
        "Pipelines are YAML: jobs, includes, variables, and workflow rules. "
        "What was missing is a lightweight way to assert organisational standards before merge.",
        size=Pt(16),
        color=BODY,
    )
    points = [
        "Focus on proving configuration does not violate standards — not end-to-end job success",
        "Open, portable alternative inspired by terraform-compliance and Conftest",
        "Complements GitLab native compliance features (higher tiers) with readable Gherkin policies",
        "Validate offline against YAML; optionally check project settings via the GitLab API",
    ]
    box = _box(s, Inches(0.9), Inches(2.55), Inches(11.5), Inches(3.8))
    _bullets(box, points, size=Pt(17))

    # 4. Two workflows
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Two workflows from one pipeline file", "From docs/index.md and usage/")
    _footer(s, 4)
    _card(s, Inches(0.65), Inches(1.7), Inches(5.85), Inches(4.7))
    _rect(s, Inches(0.65), Inches(1.7), Inches(5.85), Inches(0.65), SEA)
    _text(_box(s, Inches(0.9), Inches(1.82), Inches(5.4), Inches(0.45)), "check — Compliance", size=Pt(18), bold=True, color=WHITE)
    _bullets(
        _box(s, Inches(0.95), Inches(2.55), Inches(5.3), Inches(3.5)),
        [
            "Author Gherkin .feature policies",
            "Point at .gitlab-ci.yml",
            "Optionally enable GitLab API checks",
            "Fail the job on violations (exit 1)",
            "Reports: console, markdown, html, mr-comment, codequality",
        ],
        size=Pt(14),
    )
    _card(s, Inches(6.85), Inches(1.7), Inches(5.85), Inches(4.7))
    _rect(s, Inches(6.85), Inches(1.7), Inches(5.85), Inches(0.65), AMBER)
    _text(_box(s, Inches(7.1), Inches(1.82), Inches(5.4), Inches(0.45)), "generate — Documentation", size=Pt(18), bold=True, color=WHITE)
    _bullets(
        _box(s, Inches(7.15), Inches(2.55), Inches(5.3), Inches(3.5)),
        [
            "Build Markdown, swagger-markdown, or HTML",
            "Exclude sections; group jobs by attribute",
            "Keep pipeline knowledge current from YAML",
            "Attach docs as CI artifacts when needed",
            "Same source of truth as compliance checks",
        ],
        size=Pt(14),
    )

    # 5. BDD model
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "How policies run", "BDD reference: entities, stash, Given / When / Then")
    _footer(s, 5)
    steps = [
        ("YAML → entities", "Pipeline jobs, includes, and settings become entities with name, properties, and line numbers."),
        ("Given", "Sets the initial stash (all jobs, includes of a type, a project setting, …)."),
        ("When", "Filters the stash. If nothing matches, the scenario is skipped — not failed."),
        ("Then", "Asserts on every remaining entity. One failure fails the scenario."),
    ]
    for i, (h, body) in enumerate(steps):
        top = Inches(1.6 + i * 1.2)
        _card(s, Inches(0.65), top, Inches(12), Inches(1.05))
        _rect(s, Inches(0.65), top, Inches(0.12), Inches(1.05), SEA if i % 2 == 0 else AMBER)
        _text(_box(s, Inches(1.0), top + Inches(0.15), Inches(11.3), Inches(0.35)), h, size=Pt(16), bold=True, color=INK)
        _text(_box(s, Inches(1.0), top + Inches(0.52), Inches(11.3), Inches(0.4)), body, size=Pt(14), color=BODY)

    # 6. Policy example
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "A policy people can read", "Example from docs and image-pinning guidance")
    _footer(s, 6)
    _text(
        _box(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.4)),
        "Idea: if a job defines an image, it must not use the :latest tag.",
        size=Pt(15),
        color=BODY,
    )
    _card(s, Inches(0.65), Inches(2.1), Inches(12), Inches(4.3))
    _rect(s, Inches(0.65), Inches(2.1), Inches(12), Inches(0.08), SEA)
    _lines(
        _box(s, Inches(1.0), Inches(2.4), Inches(11.3), Inches(3.7)),
        [
            ("# METADATA", MUTED, False, Pt(15)),
            ("# title: Disallow latest image tags", MUTED, False, Pt(15)),
            ("# custom:", MUTED, False, Pt(15)),
            ("#   id: GLCI-IMAGE-PINNING-001", MUTED, False, Pt(15)),
            ("#   severity: HIGH", MUTED, False, Pt(15)),
            ("Scenario: Job images must not use the latest tag", INK, True, Pt(16)),
            ("  Given I have any job defined", SEA, True, Pt(16)),
            ("  When it has image", SEA, True, Pt(16)),
            ('  Then its image must not match ":latest$"', SEA, True, Pt(16)),
        ],
        mono=True,
    )

    # 7. Security pack
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Starter security policy pack", "examples/example-policies/security/")
    _footer(s, 7)
    policies = [
        ("Image pinning", "No :latest; prefer sha256 digests"),
        ("Component pinning", "No unpinned @main components"),
        ("Fragment pinning", "Pinned template ref: values"),
        ("Include versions", "Valid / current include refs"),
        ("Service pinning", "Versioned service images"),
        ("Execution policy", "Jobs need rules / guards"),
        ("Template extends", "Jobs follow org templates"),
        ("API hardening", "Masked vars, safer project settings"),
        ("Rules & outlines", "Advanced Scenario Outline matrices"),
    ]
    for i, (h, body) in enumerate(policies):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.6 + row * 1.65)
        _card(s, left, top, Inches(4.0), Inches(1.45))
        _rect(s, left, top, Inches(4.0), Inches(0.08), SEA)
        _text(_box(s, left + Inches(0.25), top + Inches(0.25), Inches(3.5), Inches(0.4)), h, size=Pt(14), bold=True, color=INK)
        _text(_box(s, left + Inches(0.25), top + Inches(0.7), Inches(3.5), Inches(0.55)), body, size=Pt(13), color=BODY)

    # 8. Install & quick start
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Install and first run", "Python 3.12+ · pip or Poetry")
    _footer(s, 8)
    cmds = [
        ("Install", "pip install gitlab-compliance"),
        ("Compliance", "gitlab-compliance check -f policies/ -p .gitlab-ci.yml"),
        ("With builtins", "gitlab-compliance check -f policies/ -p .gitlab-ci.yml --with-builtin"),
        ("Docs", "gitlab-compliance generate -i .gitlab-ci.yml --format swagger-markdown -o pipeline-reference.md"),
        ("Starter policies", "cp -r examples/example-policies/security/ policies/"),
    ]
    for i, (label, cmd) in enumerate(cmds):
        top = Inches(1.55 + i * 0.95)
        _card(s, Inches(0.65), top, Inches(12), Inches(0.85))
        _rect(s, Inches(0.65), top, Inches(2.3), Inches(0.85), SEA)
        _text(_box(s, Inches(0.75), top + Inches(0.25), Inches(2.1), Inches(0.4)), label, size=Pt(14), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _text(_box(s, Inches(3.2), top + Inches(0.25), Inches(9.1), Inches(0.45)), cmd, size=Pt(14), color=BODY)

    # 9. Reports
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Compliance report formats", "check --format options from usage docs")
    _footer(s, 9)
    formats = [
        ("console", "Rich terminal tables — default for local runs"),
        ("markdown", "Human-readable report file for reviews and archives"),
        ("html", "Shareable HTML compliance report"),
        ("mr-comment", "Merge request comment body; optional --post-mr-comment"),
        ("codequality", "GitLab Code Quality JSON for MR widgets"),
    ]
    for i, (name, desc) in enumerate(formats):
        top = Inches(1.55 + i * 0.95)
        _card(s, Inches(0.65), top, Inches(12), Inches(0.85))
        _text(_box(s, Inches(0.95), top + Inches(0.12), Inches(3.2), Inches(0.35)), name, size=Pt(16), bold=True, color=SEA)
        _text(_box(s, Inches(0.95), top + Inches(0.45), Inches(11.3), Inches(0.3)), desc, size=Pt(14), color=BODY)

    # 10. CI/CD patterns
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Using in CI/CD", "From docs/ci-cd/")
    _footer(s, 10)
    patterns = [
        ("Project job", "Add a compliance stage on MRs and the default branch."),
        ("Shared templates", ".compliance:offline / :api / :codequality / :oci job extends."),
        ("Central include", "Versioned jobs + policies from a security project."),
        ("GitHub Actions", "pip or container jobs in example-github-actions/."),
        ("Pipeline Execution Policy", "Org-wide inject_policy for Ultimate groups."),
        ("Pre-commit / local", "Run checks before push for fast feedback."),
    ]
    for i, (h, body) in enumerate(patterns):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.65 + row * 2.35)
        _card(s, left, top, Inches(4.0), Inches(2.15))
        _rect(s, left, top, Inches(0.1), Inches(2.15), SEA)
        _text(_box(s, left + Inches(0.3), top + Inches(0.3), Inches(3.5), Inches(0.5)), h, size=Pt(15), bold=True, color=INK)
        _text(_box(s, left + Inches(0.3), top + Inches(0.9), Inches(3.5), Inches(0.95)), body, size=Pt(13), color=BODY)

    # 11. OCI & SoD
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Policy distribution & segregation of duty", "OCI packs, metadata catalogs, central ownership")
    _footer(s, 11)
    left_items = [
        "Publish: gitlab-compliance policies push -f policies/ registry…:1.0.0",
        "Consume: check -f oci://… --update",
        "Catalog: policies doc -f policies/ -o COMPLIANCE-POLICIES.md",
        "Metadata IDs and severity feed Code Quality reports",
    ]
    right_items = [
        "Keep policy packs in a separate security repository",
        "Version refs on includes and OCI tags",
        "Authenticate private registries in CI before_script",
        "Same pack works across many consumer projects",
    ]
    _card(s, Inches(0.65), Inches(1.65), Inches(5.9), Inches(4.7))
    _text(_box(s, Inches(0.95), Inches(1.9), Inches(5.3), Inches(0.4)), "OCI policy packs", size=Pt(18), bold=True, color=INK)
    _bullets(_box(s, Inches(0.95), Inches(2.5), Inches(5.3), Inches(3.5)), left_items, size=Pt(14))
    _card(s, Inches(6.8), Inches(1.65), Inches(5.9), Inches(4.7))
    _text(_box(s, Inches(7.1), Inches(1.9), Inches(5.3), Inches(0.4)), "Central ownership", size=Pt(18), bold=True, color=INK)
    _bullets(_box(s, Inches(7.1), Inches(2.5), Inches(5.3), Inches(3.5)), right_items, size=Pt(14))

    # 12. Rollout
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Documented rollout strategy", "examples/index.md")
    _footer(s, 12)
    stages = [
        ("1", "Warn-only", "Run in CI with allow_failure: true so findings surface without blocking."),
        ("2", "Block", "Remove allow_failure once baselines are fixed; gate merge requests."),
        ("3", "Central policies", "Version policies in a security repo or OCI registry."),
        ("4", "Org-wide", "Inject compliance via Pipeline Execution Policy."),
    ]
    for i, (num, title, desc) in enumerate(stages):
        top = Inches(1.7 + i * 1.2)
        _card(s, Inches(0.65), top, Inches(12), Inches(1.05))
        _rect(s, Inches(0.8), top + Inches(0.2), Inches(0.65), Inches(0.65), SEA)
        _text(_box(s, Inches(0.8), top + Inches(0.32), Inches(0.65), Inches(0.45)), num, size=Pt(18), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _text(_box(s, Inches(1.75), top + Inches(0.15), Inches(10.5), Inches(0.35)), title, size=Pt(16), bold=True, color=INK)
        _text(_box(s, Inches(1.75), top + Inches(0.55), Inches(10.5), Inches(0.35)), desc, size=Pt(14), color=BODY)

    # 13. Requirements & commands map
    s = prs.slides.add_slide(blank)
    _bg(s)
    _header(s, "Command map & requirements", "From usage and installation docs")
    _footer(s, 13)
    cmds = [
        ("check", "Run Gherkin policies against pipeline YAML"),
        ("generate", "Build Markdown / HTML pipeline docs"),
        ("get-attributes", "Export selected job attributes as a table"),
        ("policies doc", "Policy catalog from # METADATA"),
        ("policies push/pull", "Publish or pull OCI policy bundles"),
        ("document gitstrings", "Render inline YAML fences into README tables"),
        ("release-notes", "Generate release notes from GitLab commits"),
    ]
    for i, (cmd, desc) in enumerate(cmds):
        top = Inches(1.5 + i * 0.7)
        _text(_box(s, Inches(0.8), top, Inches(3.6), Inches(0.4)), cmd, size=Pt(14), bold=True, color=SEA)
        _text(_box(s, Inches(4.5), top, Inches(8), Inches(0.4)), desc, size=Pt(14), color=BODY)
    note = _box(s, Inches(0.8), Inches(6.4), Inches(11.5), Inches(0.4))
    _text(note, "Requirements: Python 3.12+ · pipeline file · optional GITLAB_TOKEN + --project/--group for API checks", size=Pt(12), color=MUTED)

    # 14. Next steps
    s = prs.slides.add_slide(blank)
    _rect(s, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, INK)
    _rect(s, Inches(0), Inches(0), Inches(0.2), SLIDE_HEIGHT, SEA)
    _text(_box(s, Inches(0.85), Inches(1.3), Inches(11.5), Inches(0.7)), "Next steps", size=Pt(34), bold=True, color=WHITE)
    actions = [
        "Read the published docs and copy the security starter policies",
        "Run check locally, then add a warn-only CI job",
        "Generate pipeline reference docs for your team",
        "Plan central policy ownership (repo or OCI) before org-wide enforce",
    ]
    box = _box(s, Inches(0.85), Inches(2.2), Inches(11.5), Inches(2.8))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(actions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{i + 1}.  {item}"
        p.space_after = Pt(12)
        for run in p.runs:
            _font(run, size=Pt(17), color=WHITE)
    links = _box(s, Inches(0.85), Inches(5.5), Inches(11.5), Inches(1.2))
    _lines(
        links,
        [
            ("https://maturitybuilder.github.io/gitlab-compliance/", SEA_SOFT, False, Pt(14)),
            ("Installation · Usage · BDD reference · Examples · CI/CD guides", RGBColor(0xA8, 0xBE, 0xC8), False, Pt(13)),
            ("pip install gitlab-compliance", WHITE, True, Pt(14)),
        ],
    )

    out = Path("/workspace/docs/presentations/gitlab-compliance-overview.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")

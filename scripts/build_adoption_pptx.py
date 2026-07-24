#!/usr/bin/env python3
"""Build the gitlab-compliance adoption PowerPoint presentation."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Layout: widescreen 16:9
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Visual direction: deep slate + teal (enterprise calm; avoid purple/cream clichés)
NAVY = RGBColor(0x0F, 0x2A, 0x3D)
SLATE = RGBColor(0x1B, 0x3A, 0x4B)
TEAL = RGBColor(0x0D, 0x8A, 0x7A)
TEAL_LIGHT = RGBColor(0x2A, 0xB5, 0xA0)
CORAL = RGBColor(0xE0, 0x7A, 0x3D)
OFF_WHITE = RGBColor(0xF7, 0xF9, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x5A, 0x6E, 0x7A)
BODY = RGBColor(0x2C, 0x3E, 0x4A)
LIGHT_LINE = RGBColor(0xD5, 0xDE, 0xE3)


def _set_run_font(run, *, size: Pt, bold: bool = False, color: RGBColor = BODY, name: str = "Calibri") -> None:
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    latin = rpr.find(qn("a:latin"))
    if latin is None:
        latin = rpr.makeelement(qn("a:latin"), {})
        rpr.insert(0, latin)
    latin.set("typeface", name)


def _add_rect(slide, left, top, width, height, fill: RGBColor) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def _add_text_box(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def _set_single_text(shape, text: str, *, size: Pt, bold: bool = False, color: RGBColor = BODY, align=PP_ALIGN.LEFT) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        _set_run_font(run, size=size, bold=bold, color=color)


def _add_bullets(shape, items: list[str], *, size: Pt = Pt(18), color: RGBColor = BODY) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.space_after = Pt(10)
        for run in p.runs:
            _set_run_font(run, size=size, bold=False, color=color)


def _add_title_bar(slide, title: str, subtitle: str | None = None) -> None:
    _add_rect(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.35), NAVY)
    _add_rect(slide, Inches(0), Inches(1.35), SLIDE_WIDTH, Inches(0.08), TEAL)
    box = _add_text_box(slide, Inches(0.7), Inches(0.28), Inches(11.8), Inches(0.6))
    _set_single_text(box, title, size=Pt(28), bold=True, color=WHITE)
    if subtitle:
        sub = _add_text_box(slide, Inches(0.7), Inches(0.82), Inches(11.8), Inches(0.4))
        _set_single_text(sub, subtitle, size=Pt(14), color=TEAL_LIGHT)


def _add_footer(slide, page: int, total: int) -> None:
    brand = _add_text_box(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.3))
    _set_single_text(brand, "gitlab-compliance  ·  MaturityBuilder", size=Pt(11), color=MUTED)
    num = _add_text_box(slide, Inches(11.2), Inches(7.05), Inches(1.5), Inches(0.3))
    _set_single_text(num, f"{page} / {total}", size=Pt(11), color=MUTED, align=PP_ALIGN.RIGHT)


def _content_bg(slide) -> None:
    _add_rect(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, OFF_WHITE)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank = prs.slide_layouts[6]
    total = 12

    # --- 1. Title ---
    s = prs.slides.add_slide(blank)
    _add_rect(s, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, NAVY)
    _add_rect(s, Inches(0), Inches(0), Inches(0.22), SLIDE_HEIGHT, TEAL)
    # subtle accent band
    _add_rect(s, Inches(0), Inches(5.85), SLIDE_WIDTH, Inches(1.65), SLATE)

    eyebrow = _add_text_box(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.4))
    _set_single_text(eyebrow, "PLATFORM ADOPTION BRIEFING", size=Pt(14), bold=True, color=TEAL_LIGHT)

    title = _add_text_box(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.2))
    _set_single_text(title, "Adopting GitLab Compliance", size=Pt(44), bold=True, color=WHITE)

    sub = _add_text_box(s, Inches(0.9), Inches(3.5), Inches(10.5), Inches(1.0))
    _set_single_text(
        sub,
        "Catch pipeline misconfigurations before merge — with readable policies teams can share.",
        size=Pt(20),
        color=RGBColor(0xC5, 0xD4, 0xDC),
    )

    meta = _add_text_box(s, Inches(0.9), Inches(6.2), Inches(11), Inches(0.9))
    tf = meta.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "MaturityBuilder  ·  gitlab-compliance"
    for run in p.runs:
        _set_run_font(run, size=Pt(16), bold=True, color=WHITE)
    p2 = tf.add_paragraph()
    p2.text = "BDD compliance testing & pipeline documentation for GitLab CI/CD"
    for run in p2.runs:
        _set_run_font(run, size=Pt(13), color=TEAL_LIGHT)

    # --- 2. Agenda ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Agenda", "What we will cover today")
    _add_footer(s, 2, total)
    items = [
        "1.  Why pipeline compliance matters",
        "2.  What gitlab-compliance delivers",
        "3.  Two workflows: check & generate",
        "4.  How adoption works in practice",
        "5.  Rollout stages from pilot to org-wide",
        "6.  Integration patterns & next steps",
    ]
    box = _add_text_box(s, Inches(1.0), Inches(1.9), Inches(10), Inches(4.5))
    _add_bullets(box, items, size=Pt(22))

    # --- 3. The challenge ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "The challenge", "CI configuration risk is easy to miss")
    _add_footer(s, 3, total)

    cards = [
        ("Floating tags", "Jobs using :latest or unpinned includes introduce unpredictable builds and supply-chain risk."),
        ("Policy drift", "Standards live in wikis and reviews — not enforced consistently across projects."),
        ("Late discovery", "Misconfigurations often surface after merge, or in production incidents."),
        ("Shared language gap", "Security and delivery teams struggle to agree on rules that developers can act on."),
    ]
    for i, (h, body) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(1.8 + row * 2.3)
        _add_rect(s, left, top, Inches(5.9), Inches(2.0), WHITE)
        _add_rect(s, left, top, Inches(0.12), Inches(2.0), CORAL if i % 2 == 0 else TEAL)
        th = _add_text_box(s, left + Inches(0.35), top + Inches(0.3), Inches(5.3), Inches(0.45))
        _set_single_text(th, h, size=Pt(18), bold=True, color=NAVY)
        tb = _add_text_box(s, left + Inches(0.35), top + Inches(0.85), Inches(5.3), Inches(0.9))
        _set_single_text(tb, body, size=Pt(14), color=BODY)

    # --- 4. What it is ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "What is gitlab-compliance?", "Open, portable BDD checks for GitLab CI")
    _add_footer(s, 4, total)

    lead = _add_text_box(s, Inches(0.7), Inches(1.7), Inches(12), Inches(0.7))
    _set_single_text(
        lead,
        "A Python CLI that runs readable Gherkin policies against .gitlab-ci.yml — before merge.",
        size=Pt(18),
        color=BODY,
    )

    points = [
        "Negative testing: prove configuration does not violate your standards",
        "Plain-language policies developers and security can share",
        "Works offline on YAML; optional GitLab API for project settings",
        "Inspired by terraform-compliance — familiar for platform teams",
        "Integrates with GitLab CI, GitHub Actions, and pre-commit hooks",
    ]
    box = _add_text_box(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(3.8))
    _add_bullets(box, points, size=Pt(18))

    # --- 5. Value ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Why adopt it?", "Business outcomes, not just tooling")
    _add_footer(s, 5, total)

    values = [
        ("Fewer surprises", "Block risky pipeline changes in merge requests, not after release."),
        ("Faster reviews", "Automated checks reduce manual YAML scrutiny and rework."),
        ("Shared ownership", "Gherkin scenarios create a common language for risk."),
        ("Segregation of duty", "Keep policies in a separate repo or OCI registry."),
        ("Clear visibility", "Console, Markdown, HTML, MR comments, and Code Quality reports."),
        ("Living documentation", "Generate pipeline reference docs from the same YAML."),
    ]
    for i, (h, body) in enumerate(values):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.75 + row * 2.4)
        _add_rect(s, left, top, Inches(4.0), Inches(2.15), WHITE)
        _add_rect(s, left, top, Inches(4.0), Inches(0.1), TEAL)
        th = _add_text_box(s, left + Inches(0.25), top + Inches(0.35), Inches(3.5), Inches(0.45))
        _set_single_text(th, h, size=Pt(16), bold=True, color=NAVY)
        tb = _add_text_box(s, left + Inches(0.25), top + Inches(0.9), Inches(3.5), Inches(1.0))
        _set_single_text(tb, body, size=Pt(13), color=BODY)

    # --- 6. Two workflows ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Two workflows from one pipeline file", "Compliance and documentation together")
    _add_footer(s, 6, total)

    # Left card - check
    _add_rect(s, Inches(0.7), Inches(1.8), Inches(5.7), Inches(4.5), WHITE)
    _add_rect(s, Inches(0.7), Inches(1.8), Inches(5.7), Inches(0.7), TEAL)
    t1 = _add_text_box(s, Inches(0.95), Inches(1.95), Inches(5.2), Inches(0.45))
    _set_single_text(t1, "check  —  Compliance", size=Pt(20), bold=True, color=WHITE)
    b1 = _add_text_box(s, Inches(0.95), Inches(2.75), Inches(5.2), Inches(3.2))
    _add_bullets(
        b1,
        [
            "Author Gherkin .feature policies",
            "Point at .gitlab-ci.yml",
            "Optionally check GitLab API settings",
            "Fail the job on violations",
            "Report: console · markdown · html · MR · codequality",
        ],
        size=Pt(15),
    )

    # Right card - generate
    _add_rect(s, Inches(6.9), Inches(1.8), Inches(5.7), Inches(4.5), WHITE)
    _add_rect(s, Inches(6.9), Inches(1.8), Inches(5.7), Inches(0.7), CORAL)
    t2 = _add_text_box(s, Inches(7.15), Inches(1.95), Inches(5.2), Inches(0.45))
    _set_single_text(t2, "generate  —  Documentation", size=Pt(20), bold=True, color=WHITE)
    b2 = _add_text_box(s, Inches(7.15), Inches(2.75), Inches(5.2), Inches(3.2))
    _add_bullets(
        b2,
        [
            "Build Markdown or HTML reference docs",
            "Filter and group jobs by stage or attribute",
            "Keep pipeline knowledge up to date",
            "Share readable CI docs with non-YAML experts",
            "Same source of truth as compliance checks",
        ],
        size=Pt(15),
    )

    # --- 7. Policy example ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Policies people can read", "Gherkin scenarios as the contract")
    _add_footer(s, 7, total)

    intro = _add_text_box(s, Inches(0.7), Inches(1.65), Inches(12), Inches(0.5))
    _set_single_text(intro, "Example: no job may use a floating latest image tag.", size=Pt(16), color=BODY)

    _add_rect(s, Inches(0.7), Inches(2.3), Inches(12), Inches(3.8), WHITE)
    _add_rect(s, Inches(0.7), Inches(2.3), Inches(12), Inches(0.08), TEAL)
    code = _add_text_box(s, Inches(1.0), Inches(2.6), Inches(11.4), Inches(3.2))
    tf = code.text_frame
    tf.clear()
    lines = [
        "Given I have any job defined",
        "When it has image",
        'Then its image must not match ":latest$"',
        "",
        "# Runs in CI against .gitlab-ci.yml (and local includes)",
        "# Merge requests cannot introduce this class of violation",
    ]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.space_after = Pt(6)
        color = TEAL if line.startswith(("Given", "When", "Then")) else MUTED
        bold = line.startswith(("Given", "When", "Then"))
        for run in p.runs:
            _set_run_font(run, size=Pt(18), bold=bold, color=color, name="Consolas")

    # --- 8. Adoption journey ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Adoption journey", "Start small, then scale with confidence")
    _add_footer(s, 8, total)

    stages = [
        ("01", "Pilot", "One team, a few high-value policies (image pinning, include versions)."),
        ("02", "Warn", "Run in CI with allow_failure: true — surface findings without blocking."),
        ("03", "Enforce", "Remove allow_failure once baselines are fixed; gate merge requests."),
        ("04", "Centralise", "Version policies in a security repo or OCI registry."),
        ("05", "Org-wide", "Inject via Pipeline Execution Policy across groups."),
    ]
    for i, (num, title, desc) in enumerate(stages):
        top = Inches(1.65 + i * 0.95)
        _add_rect(s, Inches(0.7), top, Inches(12), Inches(0.85), WHITE)
        _add_rect(s, Inches(0.85), top + Inches(0.15), Inches(0.7), Inches(0.55), TEAL)
        nb = _add_text_box(s, Inches(0.85), top + Inches(0.22), Inches(0.7), Inches(0.45))
        _set_single_text(nb, num, size=Pt(14), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        th = _add_text_box(s, Inches(1.8), top + Inches(0.12), Inches(3), Inches(0.35))
        _set_single_text(th, title, size=Pt(16), bold=True, color=NAVY)
        tb = _add_text_box(s, Inches(1.8), top + Inches(0.45), Inches(10.5), Inches(0.35))
        _set_single_text(tb, desc, size=Pt(13), color=BODY)

    # --- 9. Integration patterns ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Where it fits in your delivery flow", "Integration patterns that stick")
    _add_footer(s, 9, total)

    patterns = [
        ("Local & pre-commit", "Developers run checks before push; fast feedback on laptop."),
        ("Project CI job", "Add a compliance job on MRs and the default branch."),
        ("Shared job templates", "Reuse .compliance:offline / :api / :codequality / :oci."),
        ("Central include", "Distribute versioned jobs + policies from a security project."),
        ("OCI policy packs", "Publish immutable bundles; consumers pull with --update."),
        ("Pipeline Execution Policy", "Org-wide injection for GitLab Ultimate groups."),
    ]
    for i, (h, body) in enumerate(patterns):
        col = i % 3
        row = i // 3
        left = Inches(0.55 + col * 4.2)
        top = Inches(1.75 + row * 2.4)
        _add_rect(s, left, top, Inches(4.0), Inches(2.15), WHITE)
        _add_rect(s, left, top, Inches(0.12), Inches(2.15), TEAL)
        th = _add_text_box(s, left + Inches(0.35), top + Inches(0.35), Inches(3.4), Inches(0.5))
        _set_single_text(th, h, size=Pt(15), bold=True, color=NAVY)
        tb = _add_text_box(s, left + Inches(0.35), top + Inches(0.95), Inches(3.4), Inches(0.9))
        _set_single_text(tb, body, size=Pt(13), color=BODY)

    # --- 10. Quick start ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "Quick start for a pilot team", "Minutes to first value")
    _add_footer(s, 10, total)

    steps = [
        ("Install", "pip install gitlab-compliance"),
        ("Copy starter policies", "cp -r examples/example-policies/security/ policies/"),
        ("Run a check", "gitlab-compliance check -f policies/ -p .gitlab-ci.yml"),
        ("Add a CI job", "Fail MRs on violations; start with allow_failure if needed"),
        ("Document pipelines", "gitlab-compliance generate -i .gitlab-ci.yml -o pipeline-reference.md"),
    ]
    for i, (title, detail) in enumerate(steps):
        top = Inches(1.7 + i * 0.95)
        _add_rect(s, Inches(0.7), top, Inches(0.9), Inches(0.75), TEAL)
        n = _add_text_box(s, Inches(0.7), top + Inches(0.18), Inches(0.9), Inches(0.45))
        _set_single_text(n, str(i + 1), size=Pt(22), bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        th = _add_text_box(s, Inches(1.9), top + Inches(0.05), Inches(10), Inches(0.35))
        _set_single_text(th, title, size=Pt(16), bold=True, color=NAVY)
        tb = _add_text_box(s, Inches(1.9), top + Inches(0.4), Inches(10.5), Inches(0.35))
        _set_single_text(tb, detail, size=Pt(14), color=MUTED)

    # --- 11. Success measures ---
    s = prs.slides.add_slide(blank)
    _content_bg(s)
    _add_title_bar(s, "How to know adoption is working", "Simple measures that matter")
    _add_footer(s, 11, total)

    measures = [
        ("Coverage", "% of active projects with compliance jobs enabled"),
        ("Prevention", "Policy violations caught in MRs vs after merge"),
        ("Time to fix", "Average time from first warning to policy-compliant YAML"),
        ("Policy maturity", "Shared central pack version in use across teams"),
        ("Clarity", "Teams can explain failed scenarios without reading CLI source"),
        ("Doc freshness", "Pipeline reference docs regenerate on each meaningful change"),
    ]
    for i, (h, body) in enumerate(measures):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 6.2)
        top = Inches(1.7 + row * 1.6)
        _add_rect(s, left, top, Inches(5.9), Inches(1.4), WHITE)
        th = _add_text_box(s, left + Inches(0.3), top + Inches(0.25), Inches(5.3), Inches(0.4))
        _set_single_text(th, h, size=Pt(16), bold=True, color=TEAL)
        tb = _add_text_box(s, left + Inches(0.3), top + Inches(0.7), Inches(5.3), Inches(0.5))
        _set_single_text(tb, body, size=Pt(14), color=BODY)

    # --- 12. Call to action ---
    s = prs.slides.add_slide(blank)
    _add_rect(s, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, NAVY)
    _add_rect(s, Inches(0), Inches(0), Inches(0.22), SLIDE_HEIGHT, TEAL)

    t = _add_text_box(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.8))
    _set_single_text(t, "Recommended next steps", size=Pt(36), bold=True, color=WHITE)

    actions = [
        "Pick one pilot project and enable warn-only compliance in CI",
        "Adopt starter security policies (image / include / component pinning)",
        "Agree a date to switch from warn to enforce",
        "Plan central policy ownership (security repo or OCI pack)",
        "Share generated pipeline docs with delivery stakeholders",
    ]
    box = _add_text_box(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(3.2))
    tf = box.text_frame
    tf.clear()
    for i, item in enumerate(actions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{i + 1}.  {item}"
        p.space_after = Pt(12)
        for run in p.runs:
            _set_run_font(run, size=Pt(18), color=WHITE)

    links = _add_text_box(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(1.0))
    tf = links.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "Docs: https://maturitybuilder.github.io/gitlab-compliance/"
    for run in p.runs:
        _set_run_font(run, size=Pt(14), color=TEAL_LIGHT)
    p2 = tf.add_paragraph()
    p2.text = "PyPI: pip install gitlab-compliance   ·   Source: github.com/MaturityBuilder/gitlab-compliance"
    for run in p2.runs:
        _set_run_font(run, size=Pt(13), color=RGBColor(0xA8, 0xBE, 0xC8))

    out = Path("/workspace/docs/presentations/gitlab-compliance-adoption.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
